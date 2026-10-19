"""tests/test_lsu_xlsx_to_md.py — xlsx→MD 转换器单元测试（合成 fixture，不触碰真实工作簿）。"""
import hashlib
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import lsu_xlsx_to_md as x2m

MAIN = x2m.MAIN_STEM


@pytest.fixture()
def fx(tmp_path):
    """合成工作簿：0 散文单列表(换行/竖线/内部空行)、1 小表、2 全宽横幅分节表、
    3 宽表(9 列→子文档, 公式无缓存, 隐藏行)、4 RunID 表(尾部全空行)。"""
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "0.说明与总览"
    ws0["A1"] = "标题行"
    ws0["A2"] = None                 # 内部空行（保留列出）
    ws0["A3"] = "多行\n第二行"
    ws0["A4"] = "竖线 | 内容"
    ws1 = wb.create_sheet("1.小表")
    ws1.append(["列A", "列B"])
    ws1.append(["x|y", 3])
    ws1.append([None, 4.5])
    ws2 = wb.create_sheet("2.分节表")   # 仿真实表 2：全宽横幅 + 二级表头
    ws2.append(["H1", "H2"])
    ws2.append(["a", "b"])
    ws2.append(["横幅文字"])
    ws2.merge_cells("A3:B3")
    ws2.append(["H1b", "H2b"])
    ws2.append(["c", "d"])
    ws3 = wb.create_sheet("3.宽表")     # 9 列 → 子文档
    ws3.append([f"列{i}" for i in range(9)])
    ws3["A2"] = "=1+1"                  # 公式：openpyxl 写入无缓存值
    ws3.append(["隐藏行内容"] + ["r"] * 8)
    ws3.row_dimensions[3].hidden = True
    ws4 = wb.create_sheet("4.矩阵")     # RunID 表（小列数 → 内联）
    ws4.append(["RunID", "值"])
    ws4.append(["A01-F0-W3", "v1"])
    ws4.append(["A01-F1-W3", "v2"])
    for _ in range(5):                  # 尾部全空行 5 行（空串物化单元格；openpyxl 对全 None 不建行）
        ws4.append(["", ""])
    p = tmp_path / "fixture.xlsx"
    wb.save(p)
    return p


# ---------- 纯函数 ----------

def test_cell_to_text():
    from datetime import datetime as dt
    assert x2m.cell_to_text(None) == ""
    assert x2m.cell_to_text(True) == "TRUE"
    assert x2m.cell_to_text(False) == "FALSE"
    assert x2m.cell_to_text(3) == "3"
    assert x2m.cell_to_text(4.5) == "4.5"
    assert x2m.cell_to_text("x") == "x"
    assert x2m.cell_to_text(dt(2026, 10, 19, 8, 30, 5)) == "2026-10-19 08:30:05"


def test_escape_unescape_roundtrip():
    for s in ["a|b", "a\\b", "x\ny", "a|b\\c\nd|e", "|", "||", "plain", "竖线 | 内容"]:
        esc = x2m.md_escape(s)
        assert "\n" not in esc
        assert "|" not in esc.replace("\\|", "")
        assert x2m.md_unescape(esc) == s


# ---------- generate ----------

def test_generate_writes_docs_and_source_untouched(fx, tmp_path):
    out = tmp_path / "out"
    sha0 = hashlib.sha256(fx.read_bytes()).hexdigest()
    report = x2m.generate(fx, out)
    assert hashlib.sha256(fx.read_bytes()).hexdigest() == sha0   # 源未被修改
    main = out / f"{MAIN}.md"
    assert main.exists()
    text = main.read_text(encoding="utf-8")
    for sec in ["## 0.说明与总览", "## 1.小表", "## 2.分节表", "## 4.矩阵"]:
        assert sec in text                                          # 内联节
    assert "## 3.宽表" not in text                                  # 宽表不在主文档
    sub = out / f"{MAIN}-03-宽表.md"
    assert sub.exists()
    assert "## 3.宽表" in sub.read_text(encoding="utf-8")
    assert report["sheet_count"] == 5
    meta = x2m.parse_metadata(text)
    assert meta["Source-SHA256"] == sha0
    assert meta["Sheet-Count"] == "5"


def test_multiline_and_pipe_rendering(fx, tmp_path):
    out = tmp_path / "out"
    x2m.generate(fx, out)
    text = (out / f"{MAIN}.md").read_text(encoding="utf-8")
    assert "多行<br>第二行" in text        # 换行 → <br>
    assert "竖线 \\| 内容" in text         # 竖线转义
    assert "x\\|y" in text


def test_banner_split(fx, tmp_path):
    out = tmp_path / "out"
    x2m.generate(fx, out)
    text = (out / f"{MAIN}.md").read_text(encoding="utf-8")
    assert x2m.BANNER_PREFIX + "横幅文字" in text   # 横幅成为分节标题
    assert "| H1b | H2b |" in text                   # 横幅后一行为新表表头
    assert "| c | d |" in text                       # 第二表数据保留
    assert "A3:B3" in text                           # 合并区间注记


def test_trailing_empty_rows_disclosed(fx, tmp_path):
    out = tmp_path / "out"
    x2m.generate(fx, out)
    text = (out / f"{MAIN}.md").read_text(encoding="utf-8")
    assert "尾部全空行 5 行" in text   # 尾部空行本身不列表（由 V3/V7 校验兜底）


def test_formula_and_hidden_notes(fx, tmp_path):
    out = tmp_path / "out"
    x2m.generate(fx, out)
    text = (out / f"{MAIN}-03-宽表.md").read_text(encoding="utf-8")
    assert "公式 A2" in text and "`=1+1`" in text and "无缓存值" in text
    assert "隐藏行 [3]" in text
    assert "| 隐藏行内容 |" in text          # 隐藏行内容仍完整列出


# ---------- validate（Task 2） ----------

def test_validate_pass_on_fresh_generation(fx, tmp_path):
    out = tmp_path / "out"
    x2m.generate(fx, out)
    assert x2m.validate(fx, out) == 0


def _mutated_fails(fx, tmp_path, mutate):
    out = tmp_path / "out"
    x2m.generate(fx, out)
    mutate(out)
    assert x2m.validate(fx, out) != 0


def _edit(out, fname, fn):
    p = out / fname
    lines = p.read_text(encoding="utf-8").splitlines()
    fn(lines)
    p.write_text("\n".join(lines) + "\n", encoding="utf-8")


def test_validate_detects_missing_row(fx, tmp_path):
    def drop(lines):
        for i, l in enumerate(lines):
            if l.startswith("| A01-F1-W3"):
                del lines[i]
                return
        raise AssertionError("fixture 行未找到")
    _mutated_fails(fx, tmp_path, lambda out: _edit(out, f"{MAIN}.md", drop))


def test_validate_detects_cell_mutation(fx, tmp_path):
    def change(lines):
        for i, l in enumerate(lines):
            if l.startswith("| A01-F0-W3"):
                lines[i] = l.replace("v1", "TAMPERED")
                return
        raise AssertionError("fixture 行未找到")
    _mutated_fails(fx, tmp_path, lambda out: _edit(out, f"{MAIN}.md", change))


def test_validate_detects_duplicate_runid(fx, tmp_path):
    def dup(lines):
        for i, l in enumerate(lines):
            if l.startswith("| A01-F1-W3"):
                lines.insert(i, l)
                return
        raise AssertionError("fixture 行未找到")
    _mutated_fails(fx, tmp_path, lambda out: _edit(out, f"{MAIN}.md", dup))


def test_validate_detects_source_drift(fx, tmp_path):
    def drift(text):
        return text.replace("Source-SHA256: ", "Source-SHA256: 0")  # 篡改绑定哈希
    def mutate(out):
        p = out / f"{MAIN}.md"
        p.write_text(drift(p.read_text(encoding="utf-8")), encoding="utf-8")
    _mutated_fails(fx, tmp_path, mutate)


def test_validate_detects_missing_sheet(fx, tmp_path):
    def drop_section(lines):
        keep, skipping = [], False
        for l in lines:
            if l.startswith('<a id="sheet-2"'):
                skipping = True
            elif l.startswith('<a id="sheet-'):
                skipping = False
            if not skipping:
                keep.append(l)
        lines[:] = keep
    _mutated_fails(fx, tmp_path, lambda out: _edit(out, f"{MAIN}.md", drop_section))


def test_validate_detects_orphan_subdoc(fx, tmp_path):
    """孤儿子文档：out_dir 多出的 MAIN_STEM-*.md 必须令 validate 失败（V1b）。"""
    def mutate(out):
        (out / f"{MAIN}-99-伪造.md").write_text("# 陈旧伪造子文档\n", encoding="utf-8")
    _mutated_fails(fx, tmp_path, mutate)


def test_validate_detects_subdoc_mutation(fx, tmp_path):
    """子文档删行必须令 validate 失败（V3 行数；钉住既有防御，43 列大表同代码路径）。"""
    def drop(lines):
        for i, l in enumerate(lines):
            if l.startswith("| 隐藏行内容"):
                del lines[i]
                return
        raise AssertionError("fixture 行未找到")
    _mutated_fails(fx, tmp_path, lambda out: _edit(out, f"{MAIN}-03-宽表.md", drop))


def test_no_trailing_blank_line(fx, tmp_path):
    """所有生成文档必须恰以单个换行结尾（git diff --check 的 blank-at-eof 必须为零）。"""
    out = tmp_path / "out"
    x2m.generate(fx, out)
    docs = sorted(out.glob("*.md"))
    assert len(docs) == 2                    # 主文档 + 1 子文档（fixture）
    for p in docs:
        data = p.read_bytes()
        assert data.endswith(b"\n"), f"{p.name} 须以换行结尾"
        assert not data.endswith(b"\n\n"), f"{p.name} 末尾不得有多余空行"
