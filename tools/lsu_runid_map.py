#!/usr/bin/env python3
"""[MDSPEC][B1] RunID 机读映射生成器（用户指令三.1，计划 2026-10-19 Task B1）。

把规格工作簿联结为每 RunID 的机读映射（规格→CLI→注入点→监控钩子→结果字段
的数据基础）：
  表 7（执行矩阵：model/unit/freq/fault_type/workload）× 表 9（传播规格：
  触发/路径/必监控/关联键/终止判定/缺失钩子/实现状态/首轮目标）×
  lsu_campaign（MODEL_BLOCKED=W6+MDSPEC 裁决；WL_BINARY/GOLDENS=实测负载
  映射，前缀匹配 + 实存校验，未映射→binary:null+注记+BLOCKED，不猜）。

用法：python3 tools/lsu_runid_map.py --xlsx <p> --out <json>
"""
import argparse
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "tools"))
import lsu_campaign as lc   # noqa: E402  (MODEL_BLOCKED / WL_BINARY / GOLDENS)

RUNID_RE = re.compile(r"^[A-Z]\d{2}-F\d+-W\d+$")


def _rows(ws):
    hdr = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in ws.iter_rows(min_row=2):
        v = [c.value for c in r]
        if v[0] is None or not str(v[0]).strip():
            continue
        out.append(v)
    return ix, out


def _txt(v):
    return "" if v is None else str(v).strip()


def _resolve_workload(spec_name):
    """WL_BINARY 前缀匹配（长键优先）+ 实存校验；未映射 → 不猜（null+注记）。

    路径解析：裸名（agu_addrmodes 等）位于 workloads/directed/；
    带子目录的值（mibench/bin/w1_all）相对 workloads/ 直取。
    """
    for key, binary in sorted(lc.WL_BINARY.items(), key=lambda kv: -len(kv[0])):
        if spec_name.startswith(key):
            for cand in (ROOT / "workloads" / "directed" / binary,
                         ROOT / "workloads" / binary):
                if cand.exists():
                    return binary, None
            return None, f"WL_BINARY 映射 {key}→{binary} 但二进制不存在"
    # WL_BINARY 未命中时查负载级结构性封锁表（README §8.3：多核 FS /
    # FS 基础设施 / SPEC 授权）——此前只输出"无映射"，掩盖了真实的外部
    # 阻塞类别（B2 triage 需要互斥原因）。长键优先，与 WL_BINARY 同序。
    for key, reason in sorted(lc.WL_BLOCKED.items(), key=lambda kv: -len(kv[0])):
        if spec_name.startswith(key):
            return None, f"workload-level:{reason}"
    return None, "规格负载名无 SE 载体映射（WL_BINARY 前缀未命中；不猜）"


def generate_map(xlsx_path):
    # [MDSPEC][R4-5.1] 模式落地登记（mode_unimpl_plan 49 项逐项落地的
    # 配套机制）：xlsx 表9「实现状态=需新增」是规格侧阻塞前提；当某模型
    # 的注入模式**已实现并实证**（真实 campaign 运行 + 注入日志非空 +
    # lsu_chain_check 逐边证据，见 mode_unimpl_plan §逐项落地模板第 5 步）
    # 后在此登记——登记后该模型的 spec 阻塞理由不再计入 status_reasons
    # （解除记录在每 entry 的 spec_landed 字段，可审计）。未登记模型保持
    # 规格阻塞原样（不得静默解除）。
    SPEC_IMPL_LANDED = {
        # C11（8 RunID，阻塞短语 "consumer-side: MSHR merge event source
        # only"）：targetField=mshr_merge（CHAOSCache 合并钩子 @
        # BaseCache::handleTimingReqMiss），campaign C11-F0-W6 3 seed
        # 3/3 activated SDC + 载体 cache_mshr_refetch ×3 全链 full——
        # 证据 docs/gem5-fi/lsu/mdspec/round4_findings.md §R4-5.1。
        # [R4-5.1-review-Minor2] 落地范围机器可读（xlsx 表9 子模型
        # C11-a..C11-d 逐一点名）：落地=C11-b 错误merge 的同 line
        # merge-目标翻位实现；C11-a（换另一合法 miss 行）/C11-c（target
        # 丢失）/C11-d（target 重复）未落地，仍为 spec 待办。
        "C11": "implemented(mshr_merge=C11-b same-line merge-target "
               "bit-flip, CHAOSCache merge hook @ "
               "BaseCache::handleTimingReqMiss; pending: C11-a cross-line "
               "substitution, C11-c target loss, C11-d target duplication; "
               "evidence: round4_findings.md §R4-5.1, campaign C11-F0-W6 "
               "3-seed + carrier x3)",
        # C12（6 RunID，阻塞短语 "consumer-side: fill-timing event source
        # only"）：targetField=mshr_fill_timing（CHAOSCache fill 推迟钩子 @
        # BaseCache::recvTimingResp 顶部，popSenderState 之前 stash + 定时
        # resume）——落地=有界时序偏移（lsuFillDelayCycles=100 缓存周期的
        # fill 交付推迟，实测偏移恰 38500 ticks = 100 个 L1D clock edge）。
        # 真机证据（round4_findings.md §R4-5.2）：campaign C12-F0-W6 3 seed
        # 3/3 activated、oracle Masked×3（时序故障数据载荷保持正确——表7
        # 钦定口径，sdc=0 是诚实结果而非缺陷）；链 2/3 terminal=full +
        # 1/3 provable_to_L0_then_unknown（seed1 被推迟行注入后未被重访，
        # L1 cache_access 缺席——按规格"缺事件只能报 Unknown"记录，不推断
        # Masked）；载体 cache_mshr_refetch ×3 checksum 全等于 golden、链
        # 3/3 full；C12-F6-W6 seed1 attempted=1/eligible=1/injected=1（无
        # 双计）链 full。[R4-5.2 同 Minor2 形态] 落地范围机器可读（xlsx
        # 表9 子模型逐一点名）：C12-a fill/writeback 丢失（无界推迟/drop）、
        # C12-b 重复、C12-c valid 过早、C12-d 旧数据+新 tag；外加触发
        # 条件列 "miss fill或dirty writeback" 的 dirty writeback 半臂
        # C12-e（有界时序偏移只落地读 miss fill 臂——资格门=非错误 READ
        # 且 sender state 为非 forward MSHR；dirty writeback 走 writeback
        # buffer 通路，不在本钩子上）。以上均未落地，仍为 spec 待办；
        # C12-e 是本仓库命名（触发条件列另一半），非表9 子模型编号。
        "C12": "implemented(mshr_fill_timing=bounded fill-delivery offset "
               "by lsuFillDelayCycles cache cycles, CHAOSCache defer hook "
               "at the top of BaseCache::recvTimingResp (stash + scheduled "
               "resume via chaosResumeDeferredResp); pending: C12-a "
               "fill/writeback drop (unbounded), C12-b duplicate fill, "
               "C12-c valid-too-early, C12-d stale-data+new-tag, C12-e "
               "dirty-writeback timing arm (the 'dirty writeback' half of "
               "trigger column 'miss fill或dirty writeback' — landed "
               "eligibility gate covers the read-miss-fill arm only; repo "
               "name for the trigger column's other half, not an xlsx "
               "sheet-9 sub-model); evidence: "
               "round4_findings.md §R4-5.2, campaign C12-F0-W6 3-seed "
               "(3/3 activated, oracle Masked x3 — timing fault keeps the "
               "data payload correct per spec; chains 2/3 full + 1/3 "
               "provable_to_L0_then_unknown) + carrier cache_mshr_refetch "
               "x3 (checksum=golden, chains 3/3 full))",
        # L04（6 RunID，阻塞短语 "consumer-side: response-pairing event
        # source only"）：--lsqfwd_l04（CHAOSLSQFwd
        # maybeSwapResponseTarget @ LSQUnit::completeDataAccess——
        # request->instruction() 配对点，09 sheet 钦定 "L04在response
        # 配对点新增可注入且可观测hook"）。
        # [Minor2 形态] 落地范围机器可读（xlsx 表9 子模型逐一点名）：
        # 落地=L04-b 同（周期）load 返回目标互换——两个在途 load 的
        # 配对表项互换，各自响应交付给对方（"两个响应数据各自合法但
        # 送错消费者"；SE 时序松弛：两半可不同周期到达，语义不变），
        # 注入行保存 source/target ID，关联=两条 seq-bound addr_registry
        # 条目（[src_va]→tgt_seq / [tgt_va]→src_seq，无任何 address-only
        # fallback——严于"标注 fallback"最低线）；未落地=L04-a 一次性
        # response transaction ID 替换（受害 load 永不完成——hang/
        # Timeout 通道，需超时簿记语义，仍为 spec 待办）。
        # 真机证据（round4_findings.md §R4-5.3）：campaign 6 cell × 3
        # seed（F0-W5 3/3 注入：Crash×2 真实 guest 页表 fault + Masked×1；
        # F0-W12/F6-W5/F6-W12 各 3/3 注入 Masked；F2-W5/W12 0 注入=诚实
        # 零——F2 平均间隔 10 万事件 >> 实测流长 324/5597，与 S01-F2-W5
        # 同构）；链：campaign 注入样本 12 个，12/12 L3→L4 proven、10/12
        # L4→L5 proven（2 个 Crash 样本截断；F2 两 cell 0 注入不计入分母）
        # （Crash 样本截断前缀无 chain_end），L0→L1/L1→L2/L2→L3 结构性
        # missing（配对点故障在 L3，受影响 load 的 L1/L2 请求路径事件
        # 先于注入——T01-postfix 同类，不回溯）。
        "L04": "implemented(resp_pairing=L04-b in-flight load response-"
               "pairing target exchange at the completeDataAccess pairing "
               "point (deferred halves: each load's returning response "
               "delivered to the OTHER load, data unchanged; SE timing "
               "relaxation documented), CHAOSLSQFwd "
               "maybeSwapResponseTarget; pending: L04-a one-shot response "
               "transaction-ID replacement (victim load never completes — "
               "hang/Timeout channel); evidence: round4_findings.md "
               "§R4-5.3, campaign L04 6-cell x 3-seed (F0/F6 12/12 "
               "injected, F0-W5 Crash x2 real guest page-table faults; "
               "F2 0-injected honest zero — mean interval 100k events >> "
               "measured streams 324/5597); chains L3->L4 proven on every "
               "injected sample, L0->L1/L1->L2/L2->L3 structurally missing "
               "(pairing-point fault is at L3; the loads' request-path "
               "events predate the injection — no back-projection))",
        # P06（6 RunID，阻塞短语 "no clean hook: prefetch-as-demand
        # marking"）：--prefetch_mode p06_as_demand（CHAOSPrefetch
        # maybeMarkFillAsDemand @ Cache::serviceMSHRTargets 的
        # setPrefetched 分类门——gem5 中 demand/prefetch 分类位的
        # 落点：from_pref 的 fill 在此成为持久的行属性
        # CacheBlk::Prefetched，替换策略 fill 数据/有用性统计/
        # hasBeenPrefetched 探测全部读它）。
        # [Minor2 形态] 落地范围机器可读（xlsx 表9 子模型逐一点名）：
        # 落地=P06-a 正向臂（把 prefetch 标成 demand）——一个
        # prefetch-owned fill 被分类为 demand（不置 Prefetched 位，
        # 该行以 demand fill 形态进入缓存；包的响应管道
        # HardPFReq/FromPrefetcher/静默 target 删除/下游
        # ReadSharedReq 全部不动——改写 cmd 会经 FromCPU 臂把
        # 未请求响应推上 CPU 侧=结构性崩溃，非合法值替换）；
        # 未落地=P06-a 反向臂（把 demand 标成 prefetch，需 demand
        # 路径分类钩子）、P06-b dirty 属性错误、P06-c ownership
        # 错误、P06-d allocate 属性错误（fill 权限属性是一致性侧
        # 状态，非预取器标记），仍为 spec 待办。
        # 证据状态（与 C11/C12/L04 不同，如实登记）：本项在禁止
        # 构建/运行 pilot 的 worktree 落地——C++ 编译未验证、
        # 真机证据 0、PF_F0_SPAN 两条跨度未实测（P06-F0 两 cell
        # 在 campaign 大声 BLOCKED(F0-span-unmeasured) 直到集成方
        # 实测）；F2/F6 四 cell 路由可解析但未运行。集成方待办见
        # R4-5.4-report.md。
        "P06": "implemented(pf_fill_class=P06-a forward arm "
               "prefetch-as-demand fill-classification marking at the "
               "Cache::serviceMSHRTargets setPrefetched gate (one "
               "prefetch-owned fill classified demand: CacheBlk::"
               "Prefetched not set, the line enters the cache as a "
               "demand fill; packet plumbing HardPFReq/FromPrefetcher/"
               "silent-target-delete untouched), CHAOSPrefetch "
               "maybeMarkFillAsDemand via fillMarkHook; pending: P06-a "
               "reverse arm (demand marked prefetch), P06-b dirty "
               "attribute, P06-c ownership attribute, P06-d allocate "
               "attribute (xlsx sheet-9 sub-models; fill-permission "
               "attributes are coherence-side state, not a prefetcher "
               "marking); evidence: NONE YET — landed compile-unverified "
               "in a no-build worktree, zero real-run samples, PF_F0_SPAN "
               "unmeasured (P06-F0 cells stay BLOCKED(F0-span-unmeasured) "
               "until integration measures; F2/F6 cells resolve runnable "
               "but unrun) — integration checklist in R4-5.4-report.md)",
        # S10（6 RunID，阻塞短语 "consumer-side: StoreSet SSID corruption
        # event source only"）：--lsqfwd_s10（CHAOSLSQFwd chaosS10* 钩子 @
        # store_set.cc StoreSet::violation 新 SSID 赋值臂 / checkInst——
        # W5 commit 7af6eed4 的 count-only 事件源由此升级为真注入）。
        # [Minor2 形态] 落地范围机器可读（xlsx 表9 子模型逐一点名）：
        # 落地=S10-a SSIT ID 换值——load 侧 SSIT 条目新赋值的 SSID 换成
        # 当前运行中另一合法在用 store-set ID（候选=SSIT 现存有效条目
        # 的 SSID 集合、排除真值；表7 实施细则"不生成任意随机位串"），
        # store 条目保留真 SSID（训练对被拆开 + load 假并入目标集合）；
        # 观测=ssit_consult（L1，checkInst 咨询腐蚀条目，before/after=
        # 反事实/实际预测 producing-store seq）与 ssit_violation（L2，
        # 腐蚀 PC 上的 ordering violation 检出）；注入行保存
        # old_ssid/new_ssid/候选数。未落地=S10-b LFST target 错误、
        # S10-c store-set valid 丢失、violation() 其余合并臂的注入点
        # （只落新 SSID 赋值臂）、多字段并发替换（表7 "每次运行只替换
        # 一个字段"已满足，maxFaults>1 的多 PC 消费面为单槽/PC）。
        # 证据状态（诚实）：落地自隔离 no-build worktree——C++ compile-
        # unverified，无 census/无 pilot/无逐边链证据；本条目解除 spec
        # "需新增"阻塞的依据是"注入模式已实现"，验证债显式登记：
        # build（scons 增量）→ F0 census（(S10,sq_forward)/(S10,gap_bfs)
        # 实测流长填 LSQFWD_F0_SPAN，F0 cell 解锁）→ campaign 6 cell ×
        # 3 seed + lsu_chain_check 逐边 + 不受影响族回归 → 再把真机
        # 证据补进本条目与 round4_findings §R4-5.7（集成侧步骤，
        # R4-5.7-report.md 清单）。
        "S10": "implemented(ssid_subst=S10-a SSIT-ID 换值 at the "
               "StoreSet::violation new-SSID assignment arm (the load's "
               "freshly assigned SSIT SSID substituted with ANOTHER LEGAL "
               "live store-set ID — candidates = valid SSIT entries' "
               "SSIDs minus the true one, workbook 07 实施细则 no-random-"
               "bit-strings; the store entry keeps the true SSID), "
               "CHAOSLSQFwd chaosS10* hooks + ssit_consult(L1)/"
               "ssit_violation(L2) observations; pending: S10-b LFST "
               "target error, S10-c store-set valid loss, the other "
               "violation() merge arms as injection points (new-SSID "
               "assignment arm only), multi-field concurrent substitution; "
               "evidence: NONE yet — landed compile-unverified from the "
               "isolated no-build worktree r4-s10, integration must build "
               "+ census LSQFWD_F0_SPAN[(S10,sq_forward)/(S10,gap_bfs)] + "
               "pilot 6 cells x 3 seeds + chain check before this entry "
               "carries real-run evidence (R4-5.7-report.md checklist))",
        # C14（2 RunID，阻塞短语 "consumer-side: MSHR busy event source
        # only"）：targetField=mshr_busy_leak（CHAOSCache maybeStickMshrBusy
        # @ BaseCache::recvTimingResp 的 MSHR 释放点）——落地=C14-a
        # "MSHR busy/occupied stuck-at-1"（已完成 fill 的表项**不释放**，
        # 留在 allocated list、永不回 free list；对应 miss 正常完成，
        # 仅槽位丢失 → 队列最终恒满 → 容量退化/Timeout）。
        # 观测=新 kind mshr_busy（L2，mshr_id/entry_index/state_before/
        # state_after/occupancy 对）+ 复用既有 mshr_* 链；inj corr=origin。
        # 证据状态（诚实）：落地自隔离 no-build worktree——**C++ compile-
        # unverified，无 census/无 pilot/无逐边链证据**；本条解除 spec
        # "需新增"的前提仅为"注入模式已实现"，验证债显式登记：集成方
        # build → F5 census（CACHE_F0_SPAN 未实测；C14 为 F5 cell 不需要
        # F0 跨度）→ campaign 2 cell × 3 seed + lsu_chain_check 逐边 +
        # 不受影响族回归 → 再把真机证据补进本条目与 round4_findings
        # §R4-5.5（R4-5.5-5.6-report.md 清单）。C14 只有一个子模型
        # （C14-a），无未落地项。
        "C14": "implemented(c14a_busy_stuck=MSHR busy/occupied stuck-at-1: "
               "the serviced MSHR entry is not released at "
               "BaseCache::recvTimingResp's release point — it stays in the "
               "queue's allocated list and never returns to the free list; "
               "the miss itself completed, only the slot is lost -> "
               "capacity degradation/Timeout), CHAOSCache "
               "maybeStickMshrBusy + mshr_busy(L2) observation; pending: "
               "NONE (single sub-model C14-a); evidence: NONE yet — landed "
               "compile-unverified from the isolated no-build worktree "
               "r4-c14c15, no census/pilot/chain evidence; integration must "
               "build + pilot 2 cells x 3 seeds + chain check before this "
               "entry carries real-run evidence (R4-5.5-5.6-report.md "
               "checklist)",
    }

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    i7, r7 = _rows(wb["7.展开执行矩阵"])
    i9, r9 = _rows(wb["9.逐RunID传播监控"])

    rid9 = {}
    for r in r9:
        rid = _txt(r[i9["RunID"]])
        if RUNID_RE.match(rid):
            if rid in rid9:
                sys.exit(f"[ABORT] 表9 RunID 重复：{rid}")
            rid9[rid] = r
    ids7 = [_txt(r[i7["RunID"]]) for r in r7 if RUNID_RE.match(_txt(r[i7["RunID"]]))]
    if set(ids7) != set(rid9):
        sys.exit(f"[ABORT] 表7∩表9 RunID 集合不一致："
                 f"only7={sorted(set(ids7) - set(rid9))[:5]} "
                 f"only9={sorted(set(rid9) - set(ids7))[:5]}")

    entries = []
    for r in r7:
        rid = _txt(r[i7["RunID"]])
        if not RUNID_RE.match(rid):
            continue
        model = _txt(r[i7["模型ID"]])
        spec_name = _txt(r[i7["负载（真实名称）"]])
        binary, wl_note = _resolve_workload(spec_name)
        s9 = rid9[rid]
        impl_status = _txt(s9[i9["实现状态"]])
        reasons = []
        if model in lc.MODEL_BLOCKED:
            reasons.append(f"campaign:{lc.MODEL_BLOCKED[model]}")
        landed = SPEC_IMPL_LANDED.get(model)
        if impl_status.startswith("需新增") and not landed:
            reasons.append(f"spec:{impl_status}")
        if wl_note:
            reasons.append(f"workload:{wl_note}")
        entries.append(dict(
            run_id=rid, model_id=model,
            unit=_txt(r[i7["单元"]]), freq=_txt(r[i7["频率"]]),
            fault_type=_txt(r[i7["故障类型"]]),
            submodels=_txt(s9[i9["故障子模型"]]),
            trigger=_txt(s9[i9["触发条件"]]),
            workload=dict(spec_name=spec_name, binary=binary,
                          golden=lc.GOLDENS.get(binary) if binary else None,
                          note=wl_note or ""),
            path_raw=_txt(s9[i9["完整传播路径（按顺序）"]]),
            path_nodes=re.findall(r"L\d", _txt(s9[i9["完整传播路径（按顺序）"]])),
            monitors=_txt(s9[i9["必监控单元/事件/指标"]]),
            existing_coverage=_txt(s9[i9["现有代码监控与源码"]]),
            missing_hooks=_txt(s9[i9["缺失监控及新增钩子"]]),
            correlation_keys=_txt(s9[i9["跨单元关联键"]]),
            termination=_txt(s9[i9["完整性/传播终止判定"]]),
            impl_status=impl_status,
            # [MDSPEC][R4-5.1] 落地登记留证（SPEC_IMPL_LANDED）：非阻塞，
            # 记录该模型 spec「需新增」前提被实现解除的依据。
            spec_landed=landed or "",
            status="BLOCKED" if reasons else "PARTIAL",
            status_reasons=reasons,
            target_activated_samples=_txt(s9[i9["首轮目标Activated样本数"]]),
        ))

    from collections import Counter
    summary = dict(total=len(entries),
                   status=dict(Counter(e["status"] for e in entries)),
                   blocked_by_unit=dict(Counter(
                       e["unit"] for e in entries if e["status"] == "BLOCKED")))
    return dict(source=Path(xlsx_path).name, entries=entries, summary=summary)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    m = generate_map(a.xlsx)
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    Path(a.out).write_text(json.dumps(m, ensure_ascii=False, indent=1),
                           encoding="utf-8")
    print(f"[runid_map] {m['summary']['total']} RunID → {a.out}")
    print(f"[runid_map] status={m['summary']['status']} "
          f"blocked_by_unit={m['summary']['blocked_by_unit']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
