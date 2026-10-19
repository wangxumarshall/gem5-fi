#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""gem5-fi LSU 方案工作簿 → Markdown 完整转换与校验.

用法:
  python3 tools/lsu_xlsx_to_md.py --xlsx <path.xlsx> --out-dir <dir>             # 生成
  python3 tools/lsu_xlsx_to_md.py --xlsx <path.xlsx> --out-dir <dir> --validate  # 校验（Task 2）

约定（与 docs/superpowers/plans/2026-10-19-lsu-xlsx-to-md.md 一致）:
  - 源 xlsx 只读（绝不 save）；生成/校验前后 sha256 必须不变。
  - canonical(v) = md_escape(cell_to_text(v).strip())；生成与校验共用，防漂移。
  - 全宽合并横幅行 → "### 分节横幅：" 分节；其后一行为新表表头。
  - 尾部全空行不列出，注记披露；内部空行照常列出；隐藏行列内容仍完整列出。
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, time as dtime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

MAIN_STEM = "gem5-fi-LSU单元故障注入方案"
SUB_DOC_MIN_COLS = 9     # max_column >= 9 → 子文档
SUB_DOC_MIN_ROWS = 101   # max_row >= 101 → 子文档
RUNID_RE = re.compile(r"^[A-Z]\d{2}-F\d+-W\d+$")
META_BEGIN = "<!-- lsu-xlsx2md-metadata"
META_END = "-->"
BANNER_PREFIX = "### 分节横幅："
ANCHOR_FMT = '<a id="sheet-{idx}"></a>'
FS_UNSAFE = re.compile(r'[\\/:*?"<>|]')


# ---------- 单元格文本化与转义（生成/校验共用，单一实现） ----------

def cell_to_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, dtime):
        return v.strftime("%H:%M:%S")
    return str(v)


def md_escape(text: str) -> str:
    text = text.replace("\\", "\\\\").replace("|", "\\|")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    return text.replace("\n", "<br>")


def md_unescape(text: str) -> str:
    text = text.replace("<br>", "\n")
    out, i, n = [], 0, len(text)
    while i < n:
        if text[i] == "\\" and i + 1 < n and text[i + 1] in "\\|":
            out.append(text[i + 1])
            i += 2
        else:
            out.append(text[i])
            i += 1
    return "".join(out)


def canonical(v) -> str:
    return md_escape(cell_to_text(v).strip())


def split_md_row(line: str) -> list:
    """把一条 MD 表格行拆回原始文本单元格（已反转义、已去边缘空白）。"""
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    cells, cur, i, n = [], [], 0, len(body)
    while i < n:
        ch = body[i]
        if ch == "\\" and i + 1 < n and body[i + 1] in "\\|":
            cur.append(body[i + 1])
            i += 2
        elif ch == "|":
            cells.append("".join(cur))
            cur = []
            i += 1
        else:
            cur.append(ch)
            i += 1
    cells.append("".join(cur))
    return [md_unescape(c.strip()) for c in cells]


# ---------- 工作表模型 ----------

@dataclass
class SheetModel:
    name: str
    ncols: int
    nrows_raw: int          # ws.max_row（含尾部全空行）
    grid: list              # 行 1..last_nonempty 的原始值矩阵（每行长度 = ncols）
    trailing_empty: int     # 被披露而不列出的尾部全空行数
    banner_rows: dict       # 行号(1-based) -> 横幅文本（全宽单行合并）
    formulas: dict          # (row, col) -> 公式串
    merged: list            # 合并区间串列表，如 ['A21:E21']
    hidden_rows: list
    hidden_cols: list       # 列号列表
    sheet_state: str        # visible / hidden / veryHidden
    nonempty: int
    runids: list            # 首列中匹配 RUNID_RE 的值（行 2 起）


def build_sheet_model(ws, wsf) -> SheetModel:
    ncols, nrows = ws.max_column, ws.max_row
    grid = [[ws.cell(row=r, column=c).value for c in range(1, ncols + 1)]
            for r in range(1, nrows + 1)]
    last_nonempty = 0
    for idx, row in enumerate(grid, 1):
        if any(v is not None and v != "" for v in row):
            last_nonempty = idx
    trailing_empty = nrows - last_nonempty
    grid = grid[:last_nonempty]
    banner_rows = {}
    merged = []
    for rng in ws.merged_cells.ranges:
        merged.append(str(rng))
        if rng.min_col == 1 and rng.max_col == ncols and rng.min_row == rng.max_row:
            banner_rows[rng.min_row] = cell_to_text(
                ws.cell(row=rng.min_row, column=1).value).strip()
    formulas = {}
    for row in wsf.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                formulas[(c.row, c.column)] = c.value
    hidden_rows = sorted(r for r, d in ws.row_dimensions.items() if d.hidden)
    hidden_cols = sorted(column_index_from_string(k)
                         for k, d in ws.column_dimensions.items() if d.hidden)
    nonempty = sum(1 for row in grid for v in row if v is not None and v != "")
    runids = [str(row[0]).strip() for row in grid[1:]
              if row and isinstance(row[0], str) and RUNID_RE.match(str(row[0]).strip())]
    return SheetModel(name=ws.title, ncols=ncols, nrows_raw=nrows, grid=grid,
                      trailing_empty=trailing_empty, banner_rows=banner_rows,
                      formulas=formulas, merged=merged, hidden_rows=hidden_rows,
                      hidden_cols=hidden_cols, sheet_state=ws.sheet_state,
                      nonempty=nonempty, runids=runids)


# ---------- 渲染 ----------

def render_table(header: list, data: list) -> list:
    lines = ["| " + " | ".join(canonical(v) for v in header) + " |",
             "|" + "---|" * len(header)]
    for row in data:
        lines.append("| " + " | ".join(canonical(v) for v in row) + " |")
    return lines


def render_sheet_section(m: SheetModel, idx: int) -> str:
    lines = [ANCHOR_FMT.format(idx=idx), f"## {m.name}"]
    runid_note = f"；RunID {len(m.runids)}（唯一）" if m.runids else ""
    state_note = "" if m.sheet_state == "visible" else f"；工作表状态={m.sheet_state}（内容仍完整转换）"
    lines.append(f"> {m.ncols} 列 × {m.nrows_raw} 行；非空单元格 {m.nonempty}{runid_note}{state_note}")
    lines.append("")
    # 按全宽横幅分节；行 1 恒为首个表的表头（12 表实测如此）
    segs = []
    header = m.grid[0] if m.grid else []
    data, expect_header, seg_banner = [], False, None
    for ridx in range(2, len(m.grid) + 1):
        row = m.grid[ridx - 1]
        if ridx in m.banner_rows:
            segs.append((seg_banner, header, data))
            seg_banner, header, data, expect_header = m.banner_rows[ridx], [], [], True
        elif expect_header:
            header, expect_header = row, False
        else:
            data.append(row)
    if m.grid:
        segs.append((seg_banner, header, data))
    for banner, header, data in segs:
        if banner is not None:
            lines.append(BANNER_PREFIX + md_escape(banner))
            lines.append("")
        if header:
            lines.extend(render_table(header, data))
            lines.append("")
    notes = []
    if m.trailing_empty:
        notes.append(f"尾部全空行 {m.trailing_empty} 行（源行 {len(m.grid) + 1}–{m.nrows_raw}）未列出，内容为空")
    if m.hidden_rows:
        notes.append(f"隐藏行 {m.hidden_rows}：内容仍完整列出")
    if m.hidden_cols:
        letters = [get_column_letter(c) for c in m.hidden_cols]
        notes.append(f"隐藏列 {letters}：内容仍完整列出")
    if m.merged:
        notes.append(f"合并单元格：{', '.join(m.merged)}（全宽横幅已作分节标题；其余仅锚单元格有值）")
    for (r, c), f in sorted(m.formulas.items()):
        disp = canonical(m.grid[r - 1][c - 1]) if r - 1 < len(m.grid) else ""
        shown = disp if disp else "（无缓存值，显示为空）"
        notes.append(f"公式 {get_column_letter(c)}{r}：`{f}` → 当前显示值：{shown}")
    if notes:
        lines.append("注：" + "；".join(notes) + "。")
        lines.append("")
    return "\n".join(lines)


def subdoc_filename(name: str, idx: int) -> str:
    m = re.match(r"^(\d+)\.(.+)$", name)
    no = f"{int(m.group(1)):02d}" if m else f"{idx:02d}"
    title = FS_UNSAFE.sub("_", m.group(2)) if m else FS_UNSAFE.sub("_", name)
    return f"{MAIN_STEM}-{no}-{title}.md"


def is_subdoc(m: SheetModel) -> bool:
    return m.ncols >= SUB_DOC_MIN_COLS or m.nrows_raw >= SUB_DOC_MIN_ROWS


def parse_metadata(text: str) -> dict:
    m = re.search(re.escape(META_BEGIN) + r"\s*\n(.*?)\n\s*" + re.escape(META_END),
                  text, re.S)
    if not m:
        return {}
    out = {}
    for line in m.group(1).splitlines():
        if ":" in line:
            k, v = line.split(":", 1)
            out[k.strip()] = v.strip()
    return out


# ---------- generate ----------

def generate(xlsx_path: Path, out_dir: Path) -> dict:
    sha = hashlib.sha256(xlsx_path.read_bytes()).hexdigest()
    stat = xlsx_path.stat()
    wbv = openpyxl.load_workbook(xlsx_path, data_only=True)   # 显示值（公式取缓存）
    wbf = openpyxl.load_workbook(xlsx_path, data_only=False)  # 原始（探公式）
    models = [build_sheet_model(wbv[name], wbf[name]) for name in wbv.sheetnames]
    # V0 前置不变式：字面 '<br>' 会让往返解析失真——响亮中止而非静默出错
    for m in models:
        for row in m.grid:
            for v in row:
                if isinstance(v, str) and "<br>" in v:
                    sys.exit(f"[ABORT] 工作表 {m.name!r} 含字面 '<br>'，"
                             f"需先扩展 md_escape/md_unescape 规则再转换")
    out_dir.mkdir(parents=True, exist_ok=True)
    toc_rows, inline_sections, written = [], [], []
    for idx, m in enumerate(models):
        if is_subdoc(m):
            fname = subdoc_filename(m.name, idx)
            text = (f"# {m.name}（{MAIN_STEM} 子表）\n\n"
                    f"> 源工作表 `{m.name}`（{m.ncols} 列 × {m.nrows_raw} 行）；"
                    f"返回[主文档]({MAIN_STEM}.md)。\n\n"
                    + render_sheet_section(m, idx))
            p = out_dir / fname
            p.write_text(text, encoding="utf-8")
            written.append(p)
            where = f"[子文档 {fname}]({fname})"
        else:
            inline_sections.append(render_sheet_section(m, idx))
            where = f"[本文 {m.name}](#sheet-{idx})"
        toc_rows.append(f"| {idx} | {m.name} | {m.ncols} 列 × {m.nrows_raw} 行 | {where} |")
    s7 = next((m for m in models if m.name == "7.展开执行矩阵"), None)
    s9 = next((m for m in models if m.name == "9.逐RunID传播监控"), None)
    meta = (f"{META_BEGIN}\n"
            f"Source-File: {xlsx_path.name}\n"
            f"Source-SHA256: {sha}\n"
            f"Source-MTime-Local: {datetime.fromtimestamp(stat.st_mtime).isoformat(timespec='seconds')}\n"
            f"Source-Size-Bytes: {stat.st_size}\n"
            f"Generated-At-Local: {datetime.now().isoformat(timespec='seconds')}\n"
            f"Sheet-Count: {len(models)}\n"
            f"RunID-Count-Sheet7: {len(s7.runids) if s7 else 'n/a'}\n"
            f"RunID-Count-Sheet9: {len(s9.runids) if s9 else 'n/a'}\n"
            f"Tool: tools/lsu_xlsx_to_md.py\n"
            f"{META_END}")
    main_doc = [f"# {MAIN_STEM}（Markdown 转换版）", "", meta, "",
                "> 本文档由上述 SHA256 对应的 xlsx 工作簿**自动完整转换**生成；源文件未被修改。",
                "> 转换规则：单元格换行→`<br>`；`|`→`\\|`；全宽合并横幅→`### 分节横幅`；"
                "尾部全空行不列出（各节注记披露）；隐藏行列内容仍完整列出。",
                "", "## 目录", "", "| # | 工作表 | 规模 | 位置 |", "|---|--------|------|------|"]
    main_doc.extend(toc_rows)
    main_doc.append("")
    main_doc.extend(inline_sections)
    p = out_dir / f"{MAIN_STEM}.md"
    p.write_text("\n".join(main_doc) + "\n", encoding="utf-8")
    written.insert(0, p)
    print(f"[generate] 源: {xlsx_path} sha256={sha[:16]}…")
    for f in written:
        print(f"[generate] 写出 {f} ({f.stat().st_size} bytes)")
    print(f"[generate] 工作表 {len(models)}；子文档 {sum(1 for m in models if is_subdoc(m))}；"
          f"RunID 表7={len(s7.runids) if s7 else 'n/a'} 表9={len(s9.runids) if s9 else 'n/a'}")
    return {"sheet_count": len(models), "written": [str(f) for f in written],
            "runid_sheet7": len(s7.runids) if s7 else None,
            "runid_sheet9": len(s9.runids) if s9 else None}


# ---------- validate（Task 2） ----------

KEY_COLUMNS = {  # V4 显式报告的关键字段（表名 → 表头名列表）
    "7.展开执行矩阵": ["RunID", "传播监控", "预期结果", "故障表现形式/子模型",
                     "Attempted", "Activated", "Masked", "Crash", "SDC",
                     "记录状态", "实测备注", "检测/告警证据"],
    "9.逐RunID传播监控": ["RunID", "完整传播路径（按顺序）", "必监控单元/事件/指标",
                       "完整性/传播终止判定", "首轮目标Activated样本数"],
    "3.位置x模型矩阵": ["模型ID"],
    "10.源码监控映射": ["模型ID"],
    "6.负载清单": ["负载ID"],
}


def parse_md_tables(text: str) -> list:
    """小节文本 → 原始文本网格（行序 = 源行序；横幅行为单格行；跳过非表格行）。"""
    grid, lines, i = [], text.splitlines(), 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            cells = [split_md_row(l) for l in block]
            if len(cells) >= 2:                      # 表头 + 分隔行 + 数据…
                grid.append(cells[0])
                grid.extend(cells[2:])
        elif line.startswith(BANNER_PREFIX):
            grid.append([md_unescape(line[len(BANNER_PREFIX):].strip())])
            i += 1
        else:
            i += 1
    return grid


def split_main_sections(main_text: str) -> dict:
    parts = re.split(r'^<a id="sheet-(\d+)"></a>$', main_text, flags=re.M)
    return {int(parts[j]): parts[j + 1] for j in range(1, len(parts), 2)}


def md_row_to_canonical(cells: list, ncols: int) -> list:
    out = [md_escape(c.strip()) for c in cells]
    if len(out) < ncols:
        out += [""] * (ncols - len(out))
    return out[:ncols]


def validate(xlsx_path: Path, out_dir: Path) -> int:
    checks, failures = [], []

    def check(name, ok):
        checks.append((name, ok))
        if not ok:
            failures.append(name)

    sha_before = hashlib.sha256(xlsx_path.read_bytes()).hexdigest()
    main_path = out_dir / f"{MAIN_STEM}.md"
    if not main_path.exists():
        print(f"[validate] 缺主文档 {main_path}")
        return 1
    main_text = main_path.read_text(encoding="utf-8")
    meta = parse_metadata(main_text)

    wbv = openpyxl.load_workbook(xlsx_path, data_only=True)
    wbf = openpyxl.load_workbook(xlsx_path, data_only=False)
    names = list(wbv.sheetnames)
    models = {n: build_sheet_model(wbv[n], wbf[n]) for n in names}

    # V0 源不变式
    lit_br = sum(1 for m in models.values() for row in m.grid for v in row
                 if isinstance(v, str) and "<br>" in v)
    edge_ws = sum(1 for m in models.values() for row in m.grid for v in row
                  if isinstance(v, str) and v != v.strip())
    check(f"V0 无字面 '<br>' 单元格（实测 {lit_br}）", lit_br == 0)
    check(f"V0 无边缘空白单元格（实测 {edge_ws}）", edge_ws == 0)
    n_formula = sum(len(m.formulas) for m in models.values())
    n_hidden_r = sum(len(m.hidden_rows) for m in models.values())
    n_hidden_c = sum(len(m.hidden_cols) for m in models.values())
    n_hidden_sh = sum(1 for m in models.values() if m.sheet_state != "visible")
    checks.append((f"V0 报告: 公式 {n_formula}；隐藏行 {n_hidden_r}；"
                   f"隐藏列 {n_hidden_c}；非 visible 工作表 {n_hidden_sh}", True))

    # V5 源绑定
    check(f"V5 元数据 Source-SHA256 == 当前源", meta.get("Source-SHA256") == sha_before)
    check("V5 元数据 Sheet-Count == 实际", meta.get("Sheet-Count") == str(len(names)))

    # 定位每个表的 MD 文本
    sections = split_main_sections(main_text)
    md_of = {}
    for idx, name in enumerate(names):
        m = models[name]
        if is_subdoc(m):
            p = out_dir / subdoc_filename(name, idx)
            if not p.exists():
                check(f"V1 [{name}] 子文档存在", False)
                continue
            check(f"V1 [{name}] 子文档存在", True)
            md_of[name] = p.read_text(encoding="utf-8")
        else:
            md_of[name] = sections.get(idx, "")
            check(f"V1 [{name}] 内联节存在", bool(md_of[name]))

    # V1 目录链接可解析
    toc_lines = [l for l in main_text.splitlines()
                 if l.startswith("| ") and ("](#" in l or "](gem5-fi-" in l)]
    n_toc = len(toc_lines)
    check(f"V1 目录行数 == 工作表数（{n_toc} vs {len(names)}）", n_toc == len(names))
    bad_links = 0
    for l in toc_lines:
        for target in re.findall(r"\]\(([^)]+)\)", l):
            if target.startswith("#"):
                if f'<a id="{target[1:]}"></a>' not in main_text:
                    bad_links += 1
            else:
                if not (out_dir / target).exists():
                    bad_links += 1
    check(f"V1 目录链接全部可解析（断链 {bad_links}）", bad_links == 0)

    # 每表 canonical 网格（xlsx 侧）
    xcan = {n: [[canonical(v) for v in row] for row in m.grid]
            for n, m in models.items()}

    # V3 全量往返 + V7 披露注记
    for name in names:
        m = models[name]
        parsed = parse_md_tables(md_of.get(name, ""))
        pcan = [md_row_to_canonical(r, m.ncols) for r in parsed]
        if len(pcan) != len(m.grid):
            check(f"V3 [{name}] 行数一致（xlsx {len(m.grid)} vs md {len(pcan)}）", False)
            continue
        diff = next((r for r, (x, p) in enumerate(zip(xcan[name], pcan), 1) if x != p), None)
        check(f"V3 [{name}] 逐格相等（{m.nonempty} 非空格）", diff is None)
        if diff:
            print(f"    V3[{name}] 首个差异行 {diff}: xlsx={xcan[name][diff-1][:3]} "
                  f"md={pcan[diff-1][:3]}")
        if m.trailing_empty:
            check(f"V7 [{name}] 尾部空行注记（{m.trailing_empty} 行）",
                  f"尾部全空行 {m.trailing_empty} 行" in md_of.get(name, ""))
        if m.merged:
            check(f"V7 [{name}] 合并单元格注记",
                  all(r in md_of.get(name, "") for r in m.merged))
        if m.formulas:
            check(f"V7 [{name}] 公式注记（{len(m.formulas)} 条）",
                  "公式 " in md_of.get(name, ""))

    # V2 RunID（对每个含 RunID 的表）
    rid_sets = {}
    for name in names:
        m = models[name]
        if not m.runids:
            continue
        parsed = parse_md_tables(md_of.get(name, ""))
        md_rids = [r[0] for r in parsed[1:] if r and RUNID_RE.match(r[0])]
        check(f"V2 [{name}] RunID 数量一致（xlsx {len(m.runids)} vs md {len(md_rids)}）",
              len(md_rids) == len(m.runids))
        check(f"V2 [{name}] RunID 无重复", len(md_rids) == len(set(md_rids)))
        check(f"V2 [{name}] RunID 集合与源一致", set(md_rids) == set(m.runids))
        rid_sets[name] = set(m.runids)
    pair = [n for n in ("7.展开执行矩阵", "9.逐RunID传播监控") if n in rid_sets]
    if len(pair) == 2:
        check("V2 表7/表9 RunID 集合相等", rid_sets[pair[0]] == rid_sets[pair[1]])
    else:
        checks.append(("V2 表7/表9 RunID 集合相等：n/a（源不足两张 RunID 表）", True))

    # V4 关键字段非空计数
    for name, wants in KEY_COLUMNS.items():
        if name not in models:
            checks.append((f"V4 [{name}] n/a（源无此表）", True))
            continue
        m, parsed = models[name], parse_md_tables(md_of.get(name, ""))
        pcan = [md_row_to_canonical(r, m.ncols) for r in parsed]
        hdr = xcan[name][0] if xcan[name] else []
        for want in wants:
            ci = next((i for i, h in enumerate(hdr) if h == canonical(want)), None)
            if ci is None:
                check(f"V4 [{name}][{want}] 找到列", False)
                continue
            n_x = sum(1 for row in xcan[name][1:] if row[ci])
            n_m = sum(1 for row in pcan[1:] if ci < len(row) and row[ci])
            check(f"V4 [{name}][{want}] 非空 xlsx={n_x} md={n_m}", n_x == n_m)

    # V6 校验期间源未被修改
    sha_after = hashlib.sha256(xlsx_path.read_bytes()).hexdigest()
    check("V6 校验前后源 sha256 一致（只读证明）", sha_after == sha_before)

    print("=" * 72)
    for name, ok in checks:
        print(f"  [{'PASS' if ok else 'FAIL'}] {name}")
    print("=" * 72)
    if failures:
        print(f"VALIDATE FAIL（{len(failures)} 项失败）:")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("ALL PASS")
    return 0


# ---------- CLI ----------

def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="gem5-fi LSU 方案工作簿 xlsx→Markdown 转换与校验")
    ap.add_argument("--xlsx", required=True, help="源工作簿路径（只读，绝不修改）")
    ap.add_argument("--out-dir", required=True, help="Markdown 输出目录")
    ap.add_argument("--validate", action="store_true", help="校验 out-dir 中已生成文档")
    a = ap.parse_args(argv)
    if a.validate:
        return validate(Path(a.xlsx), Path(a.out_dir))
    generate(Path(a.xlsx), Path(a.out_dir))
    return 0


if __name__ == "__main__":
    sys.exit(main())
